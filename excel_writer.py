from datetime import date, datetime
from pathlib import Path
import os
import tempfile
from typing import Any

from openpyxl import load_workbook


SHEET_SPECS = {
    "gas.day_ahead": ("GAZ_RDN", ("price_pln_mwh", "volume_mwh", "change_pct")),
    "gas.futures": ("GAZ_OTF", ("dkr_pln_mwh", "max_pln_mwh", "volume_mwh")),
    "electricity.day_ahead": ("EE_RDN", ("price_pln_mwh", "volume_mwh", "change_pct")),
    "electricity.futures": ("EE_OTF", ("dkr_pln_mwh", "max_pln_mwh", "volume_mwh")),
    "coal.futures": ("WEGIEL", ("price", "open", "high", "low", "volume", "change_pct")),
    "emissions.futures": ("EUA", ("price", "open", "high", "low", "volume", "change_pct")),
}


def _records(payload: dict[str, Any], section: str) -> list[dict[str, Any]]:
    market, collection = section.split(".")
    return payload["data"][market][collection]


def _upsert_sheet(
    worksheet,
    records: list[dict[str, Any]],
    value_fields: tuple[str, ...],
    retrieved_at: str,
) -> None:
    headers = [cell.value for cell in worksheet[1]]
    expected = ["trading_date", "data_id", "source_url", "instrument", *value_fields, "retrieved_at"]
    if headers != expected:
        raise ValueError(
            f"Arkusz {worksheet.title} ma nieoczekiwane kolumny: {headers!r}."
        )
    existing: dict[tuple[str, str], int] = {}
    last_data_row = 1
    for row in range(2, worksheet.max_row + 1):
        trading_date = worksheet.cell(row, 1).value
        if hasattr(trading_date, "isoformat"):
            trading_date = trading_date.isoformat()[:10]
        data_id = worksheet.cell(row, 2).value
        if trading_date and data_id:
            existing[(str(trading_date)[:10], str(data_id))] = row
            last_data_row = max(last_data_row, row)
    for record in records:
        key = (record["trading_date"], record["data_id"])
        row = existing.get(key, last_data_row + 1)
        values = record["values"]
        output = [
            record["trading_date"], record["data_id"], record["source_url"],
            record["instrument"], *(values[field] for field in value_fields), retrieved_at,
        ]
        for column, value in enumerate(output, 1):
            worksheet.cell(row, column, value)
        existing[key] = row
        last_data_row = max(last_data_row, row)


def _delete_window(worksheet, date_window: tuple[date, date]) -> None:
    start, end = date_window
    for row in range(worksheet.max_row, 1, -1):
        value = worksheet.cell(row, 1).value
        try:
            parsed = value.date() if isinstance(value, datetime) else (
                value if isinstance(value, date) else date.fromisoformat(str(value)[:10])
            )
        except (TypeError, ValueError):
            continue
        if start <= parsed <= end:
            worksheet.delete_rows(row)


def save_market_data(
    workbook_path: Path,
    payload: dict[str, Any],
    coal_window: tuple[date, date],
    eua_window: tuple[date, date],
) -> None:
    workbook = load_workbook(workbook_path)
    temp_path: Path | None = None
    try:
        retrieved_at = payload.get("retrieved_at") or datetime.now().astimezone().isoformat()
        for section, (sheet_name, fields) in SHEET_SPECS.items():
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Brak arkusza {sheet_name} w DataMaster.xlsx.")
            worksheet = workbook[sheet_name]
            if section == "coal.futures":
                _delete_window(worksheet, coal_window)
            elif section == "emissions.futures":
                _delete_window(worksheet, eua_window)
            _upsert_sheet(worksheet, _records(payload, section), fields, retrieved_at)
        with tempfile.NamedTemporaryFile(
            prefix="DataMaster_", suffix=".xlsx", dir=workbook_path.parent, delete=False
        ) as temporary:
            temp_path = Path(temporary.name)
        workbook.save(temp_path)
        check = load_workbook(temp_path, read_only=True, data_only=True)
        check.close()
        os.replace(temp_path, workbook_path)
        temp_path = None
    finally:
        workbook.close()
        if temp_path is not None and temp_path.exists():
            temp_path.unlink()
