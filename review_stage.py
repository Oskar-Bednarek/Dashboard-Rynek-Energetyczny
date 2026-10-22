from datetime import date
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from excel_writer import SHEET_SPECS, save_market_data
from json_validator import ValidationReport, validate_response
from manual_completion import apply_manual_records, find_missing_tge_records


def flatten_market_records(payload: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    data = payload.get("data", {})
    if not isinstance(data, dict):
        return rows
    for market, market_data in data.items():
        if not isinstance(market_data, dict):
            continue
        for collection, records in market_data.items():
            if not isinstance(records, list):
                continue
            for record in records:
                if not isinstance(record, dict):
                    continue
                row = {
                    "sekcja": f"{market}.{collection}",
                    "data": record.get("trading_date"),
                    "instrument": record.get("instrument"),
                    "data_id": record.get("data_id"),
                    "source_url": record.get("source_url"),
                }
                values = record.get("values")
                if isinstance(values, dict):
                    row.update(values)
                rows.append(row)
    return rows


def build_dashboard_preview(
    workbook_path: Path,
    payload: dict[str, Any],
    history_window: tuple[date, date],
    otf_session_date: date,
) -> list[dict[str, Any]]:
    combined: dict[tuple[str, str, str], dict[str, Any]] = {}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        for section, (sheet_name, value_fields) in SHEET_SPECS.items():
            worksheet = workbook[sheet_name]
            headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
            for values in worksheet.iter_rows(min_row=2, values_only=True):
                record = dict(zip(headers, values))
                if not record.get("trading_date") or not record.get("data_id"):
                    continue
                text_date = str(record["trading_date"])[:10]
                row = {
                    "sekcja": section,
                    "data": text_date,
                    "instrument": record.get("instrument"),
                    "data_id": record.get("data_id"),
                    "source_url": record.get("source_url"),
                }
                row.update({field: record.get(field) for field in value_fields})
                combined[(section, text_date, str(record["data_id"]))] = row
    finally:
        workbook.close()

    for row in flatten_market_records(payload):
        combined[(row["sekcja"], str(row["data"]), str(row["data_id"]))] = row

    start, end = history_window
    result = []
    for row in combined.values():
        try:
            row_date = date.fromisoformat(str(row["data"])[:10])
        except ValueError:
            continue
        if row["sekcja"] in {"gas.futures", "electricity.futures"}:
            if row_date != otf_session_date:
                continue
        elif not start <= row_date <= end:
            continue
        result.append(row)
    return sorted(result, key=lambda row: (str(row["sekcja"]), str(row["data"]), str(row["data_id"])))


def _source_errors(report: ValidationReport) -> list[dict[str, str]]:
    if report.data is None:
        return []
    result = []
    for error in report.data.get("errors", []):
        if isinstance(error, dict):
            result.append({
                "Sekcja": str(error.get("section", "")),
                "Link": str(error.get("source_url", "")),
                "Problem": str(error.get("message", "")),
            })
    return result


def show_review_page(workbook_path: Path) -> None:
    import streamlit as st

    report: ValidationReport | None = st.session_state.get("validation_report")
    if report is None or report.data is None or not report.is_structurally_valid:
        st.session_state["page"] = "prompt"
        st.rerun()

    source_dates: dict[str, list[date]] = st.session_state["confirmed_source_dates"]
    otf_session_date: date = st.session_state["otf_session_date"]
    coal_window: tuple[date, date] = st.session_state["coal_window"]
    eua_window: tuple[date, date] = st.session_state["eua_window"]

    st.title("Kontrola i uzupełnianie danych")
    st.success("Struktura JSON-a jest poprawna.")

    missing_records = find_missing_tge_records(
        report.data, source_dates, otf_session_date
    )
    source_errors = _source_errors(report)

    st.subheader("Dane wymagające uzupełnienia")
    if not missing_records and not source_errors:
        st.success("Nie wykryto brakujących danych ani błędów źródeł.")
    else:
        if source_errors:
            st.write(
                "Poniższe problemy zostały zgłoszone przez odpowiedź. "
                "Dla węgla i EUA sprawdź źródło przed zapisem."
            )
            st.dataframe(
                source_errors,
                hide_index=True,
                width="stretch",
                column_config={"Link": st.column_config.LinkColumn("Link")},
            )

        if missing_records:
            st.write(
                "Wpisz wartości z podanych stron. Puste pole oznacza świadomie "
                "potwierdzony brak wartości (null)."
            )
            completed = []
            with st.form("review_manual_completion"):
                for index, missing in enumerate(missing_records):
                    st.markdown(
                        f"**{missing.instrument} — {missing.trading_date}**  "
                        f"[{missing.source_url}]({missing.source_url})"
                    )
                    columns = st.columns(len(missing.value_fields) + 1)
                    values = {}
                    for column, field in zip(columns, missing.value_fields):
                        values[field] = column.number_input(
                            field,
                            value=None,
                            key=f"review_{index}_{field}",
                        )
                    confirmed = columns[-1].checkbox(
                        "Potwierdzono", key=f"review_{index}_confirmed"
                    )
                    completed.append((missing, values, confirmed))
                apply_manual = st.form_submit_button(
                    "Zastosuj ręczne uzupełnienia", type="primary"
                )
            if apply_manual:
                if not all(item[2] for item in completed):
                    st.error("Potwierdź wszystkie widoczne brakujące rekordy.")
                else:
                    updated = apply_manual_records(
                        report.data,
                        [(missing, values) for missing, values, _ in completed],
                    )
                    updated_text = json.dumps(updated, ensure_ascii=False)
                    updated_report = validate_response(
                        updated_text,
                        source_dates,
                        otf_session_date,
                        coal_window,
                        eua_window,
                    )
                    st.session_state["response"] = updated_text
                    st.session_state["validation_report"] = updated_report
                    st.session_state["parsed_response"] = updated
                    st.session_state["validated_payload"] = updated
                    st.session_state["validated_sections"] = dict(
                        updated_report.section_validity
                    )
                    st.rerun()

    st.subheader("Wszystkie dane przeznaczone do zapisu i dashboardu")
    rows = build_dashboard_preview(
        workbook_path, report.data, coal_window, otf_session_date
    )
    if rows:
        display_rows = [
            {
                key: (str(value) if key == "change_pct" and value is not None else value)
                for key, value in row.items()
            }
            for row in rows
        ]
        st.dataframe(
            display_rows,
            hide_index=True,
            width="stretch",
            column_config={"source_url": st.column_config.LinkColumn("source_url")},
        )
    else:
        st.info("JSON nie zawiera jeszcze żadnych rekordów danych.")

    st.subheader("Zapis danych")
    if report.is_valid:
        if st.button("Zapisz dane do DataMaster.xlsx", type="primary"):
            try:
                save_market_data(
                    workbook_path, report.data, coal_window, eua_window
                )
            except (OSError, ValueError, KeyError) as error:
                st.error(f"Nie udało się zapisać danych: {error}")
            else:
                st.session_state["data_saved"] = True
                st.session_state.pop("date_plan", None)
                st.success("Dane zostały zapisane w DataMaster.xlsx.")
    else:
        invalid_sections = [
            section for section, valid in report.section_validity.items() if not valid
        ]
        st.warning(
            "Zapis jest zablokowany do czasu uzupełnienia danych. "
            + (
                "Sekcje wymagające poprawy: " + ", ".join(invalid_sections)
                if invalid_sections
                else "Pozostały błędy opisane w raporcie."
            )
        )
        st.button(
            "Zapisz dane do DataMaster.xlsx",
            disabled=True,
            type="primary",
            key="review_save_disabled",
        )

    left, right = st.columns(2)
    if left.button("Wróć do promptu"):
        st.session_state["page"] = "prompt"
        st.rerun()
    if right.button("Wróć do wyboru dat"):
        st.session_state["page"] = "dates"
        st.rerun()
