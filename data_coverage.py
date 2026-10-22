from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook

@dataclass(frozen=True)
class CoverageResult:
    complete: bool
    details: dict[str, str]

def _date_text(value):
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    if isinstance(value, str) and value: return value[:10]
    return None

def assess_coverage(workbook_path: Path, rdn_missing, otf_session_date, coal_window, eua_window):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    details = {}
    try:
        for key in ("gas_rdn", "electricity_rdn"):
            count = len(rdn_missing[key])
            details[key] = "komplet" if count == 0 else f"brakuje {count} dat"
        for key, sheet_name in (("gas_otf", "GAZ_OTF"), ("electricity_otf", "EE_OTF")):
            ws = workbook[sheet_name]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            di, ii = headers.index("trading_date"), headers.index("data_id")
            ids = {str(r[ii]) for r in ws.iter_rows(min_row=2, values_only=True) if _date_text(r[di]) == otf_session_date.isoformat() and r[ii]}
            details[key] = "komplet" if len(ids) == 10 else f"jest {len(ids)}/10 kontraktów"
        for key, sheet_name, window in (("coal", "WEGIEL", coal_window), ("emissions", "EUA", eua_window)):
            ws = workbook[sheet_name]
            headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            di = headers.index("trading_date")
            dates = set()
            for row in ws.iter_rows(min_row=2, values_only=True):
                text = _date_text(row[di])
                if text:
                    parsed = date.fromisoformat(text)
                    if window[0] <= parsed <= window[1]: dates.add(parsed)
            covered = len(dates) >= 15 and min(dates, default=window[1]) <= window[0] + timedelta(days=4) and max(dates, default=window[0]) >= window[1] - timedelta(days=4)
            details[key] = "komplet" if covered else f"niewystarczające pokrycie ({len(dates)} sesji)"
    finally:
        workbook.close()
    return CoverageResult(all(v == "komplet" for v in details.values()), details)
