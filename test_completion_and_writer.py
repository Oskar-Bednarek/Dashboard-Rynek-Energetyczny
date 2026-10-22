from copy import deepcopy
from datetime import date
from pathlib import Path
import shutil
import tempfile
import unittest

from openpyxl import load_workbook

from excel_writer import save_market_data
from manual_completion import apply_manual_records, find_missing_tge_records
from review_stage import build_dashboard_preview, flatten_market_records
from test_json_validator import SESSION_DATE, SOURCE_DATES, WINDOW, valid_payload


class ManualCompletionTests(unittest.TestCase):
    def test_missing_rdn_record_can_be_completed(self) -> None:
        payload = valid_payload()
        payload["data"]["gas"]["day_ahead"] = []
        payload["status"] = "partial"
        payload["errors"] = [{
            "section": "gas.day_ahead",
            "source_url": "https://tge.pl/gaz-rdn?dateShow=20-10-2026",
            "message": "Brak rekordu.",
        }]
        missing = find_missing_tge_records(payload, SOURCE_DATES, SESSION_DATE)
        gas_missing = [item for item in missing if item.section == "gas.day_ahead"]
        self.assertEqual(len(gas_missing), 1)
        updated = apply_manual_records(
            payload,
            [(gas_missing[0], {
                "price_pln_mwh": 100.0,
                "volume_mwh": 20.0,
                "change_pct": None,
            })],
        )
        self.assertEqual(len(updated["data"]["gas"]["day_ahead"]), 1)
        self.assertEqual(updated["status"], "success")

    def test_all_missing_tge_records_are_still_discoverable(self) -> None:
        payload = valid_payload()
        payload["data"]["gas"]["day_ahead"] = []
        payload["data"]["gas"]["futures"] = []
        payload["data"]["electricity"]["day_ahead"] = []
        payload["data"]["electricity"]["futures"] = []
        missing = find_missing_tge_records(payload, SOURCE_DATES, SESSION_DATE)
        self.assertEqual(len(missing), 25)

    def test_review_table_contains_all_market_records(self) -> None:
        rows = flatten_market_records(valid_payload())
        self.assertGreaterEqual(len(rows), 27)
        keys = {(row["sekcja"], row["data"], row["data_id"]) for row in rows}
        self.assertEqual(len(keys), len(rows))
        self.assertTrue(any(row["sekcja"] == "coal.futures" for row in rows))
        self.assertTrue(any(row["instrument"] == "TGEgasDA" for row in rows))

    def test_dashboard_preview_merges_workbook_and_payload(self) -> None:
        rows = build_dashboard_preview(
            Path("DataMaster.xlsx"), valid_payload(), WINDOW, SESSION_DATE
        )
        self.assertGreaterEqual(len(rows), 27)
        keys = {(row["sekcja"], row["data"], row["data_id"]) for row in rows}
        self.assertEqual(len(keys), len(rows))


class ExcelWriterTests(unittest.TestCase):
    def test_complete_payload_is_written_to_copy(self) -> None:
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            target = Path(directory) / "DataMaster.xlsx"
            shutil.copy2("DataMaster.xlsx", target)
            payload = deepcopy(valid_payload())
            save_market_data(target, payload, WINDOW, WINDOW)
            workbook = load_workbook(target, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook["GAZ_RDN"]["B2"].value, "tge_gas_da")
                self.assertEqual(
                    workbook["EE_OTF"]["B2"].value,
                    "power_base_week_1",
                )
                self.assertEqual(
                    workbook["WEGIEL"]["B2"].value,
                    "coal_rotterdam_futures",
                )
                self.assertEqual(workbook["EUA"]["B2"].value, "eua_year_futures")
            finally:
                workbook.close()

    def test_otf_history_is_preserved_for_next_session(self) -> None:
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as directory:
            target = Path(directory) / "DataMaster.xlsx"
            shutil.copy2("DataMaster.xlsx", target)
            first = deepcopy(valid_payload())
            save_market_data(target, first, WINDOW, WINDOW)
            second = deepcopy(valid_payload())
            for market in ("gas", "electricity"):
                for record in second["data"][market]["futures"]:
                    record["trading_date"] = "2026-10-22"
                    record["source_url"] = record["source_url"].replace(
                        "21-10-2026", "22-10-2026"
                    )
            save_market_data(target, second, WINDOW, WINDOW)
            workbook = load_workbook(target, read_only=True, data_only=True)
            try:
                for sheet_name in ("GAZ_OTF", "EE_OTF"):
                    rows = [
                        row for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True)
                        if row[0] in {"2026-10-21", "2026-10-22"}
                    ]
                    self.assertEqual(len(rows), 20)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
